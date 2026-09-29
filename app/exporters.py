"""Report exporters: Markdown, DOCX (python-docx), PDF (fpdf2 with a bundled Unicode font).

Both are pure Python, so exports work on any free host without Cairo/Pango system libraries.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

import markdown as md
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fpdf import FPDF, FontFace, TextStyle

FONT_DIR = Path(__file__).parent / "fonts"
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|~~[^~]+~~|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`|\[[^\]]+\]\([^)\s]+\))")


def safe_filename(title: str) -> str:
    s = re.sub(r"[^\w\s-]", "", title).strip().lower()
    return re.sub(r"[\s_-]+", "-", s)[:60] or "research-report"


def _footer_note(report: dict) -> str:
    gen = str(report.get("generated_at", ""))[:10] or datetime.utcnow().strftime("%Y-%m-%d")
    return f"Research question: {report.get('question', '')}  ·  Generated {gen} by Agentic Research"


# ---------------------------------------------------------------- DOCX

def _add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FBF")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _add_inline(paragraph, text: str) -> None:
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("~~") and tok.endswith("~~"):
            paragraph.add_run(tok[2:-2]).font.strike = True
        elif (tok.startswith("*") and tok.endswith("*")) or (tok.startswith("_") and tok.endswith("_") and len(tok) > 2):
            paragraph.add_run(tok[1:-1]).italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
        elif tok.startswith("[") and "](" in tok:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok).groups()
            _add_hyperlink(paragraph, url, label)
        else:
            run = paragraph.add_run(tok)
            if re.fullmatch(r"(\[S\d+\])+", tok.strip()):
                run.font.color.rgb = RGBColor(0x1F, 0x5F, 0xBF)


def _add_table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    ncols = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=ncols)
    table.style = "Light Grid Accent 1"
    for i, r in enumerate(cells):
        for j in range(ncols):
            p = table.cell(i, j).paragraphs[0]
            _add_inline(p, r[j] if j < len(r) else "")
            if i == 0:
                for run in p.runs:
                    run.bold = True


def to_docx(report: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = report["markdown"].splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or re.fullmatch(r"[-*_]{3,}", s):
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=0 if level == 1 else min(level - 1, 4))
            _add_inline(h, m.group(2))
            if level == 1:
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = sub.add_run(_footer_note(report))
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif re.match(r"^\s*[-*+]\s+", line):
            indent = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
            _add_inline(p, re.sub(r"^\s*[-*+]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _add_inline(p, s.lstrip("> "))
        else:
            _add_inline(doc.add_paragraph(), s)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

class _PDF(FPDF):
    footer_text = ""

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("DejaVu", "I", 7.5)
        self.set_text_color(120, 120, 120)
        self.cell(0, 6, f"{self.footer_text[:120]}   ·   Page {self.page_no()}/{{nb}}", align="C")


def to_pdf(report: dict) -> bytes:
    text = report["markdown"]
    text = re.sub(r"~~([^~]+)~~", r"<i>[removed] \1</i>", text)
    text = re.sub(r"(\[S\d+\])", r'<font color="#1F5FBF">\1</font>', text)
    html = md.markdown(text, extensions=["tables", "sane_lists"])
    html = html.replace("<hr />", "<br>")
    # fpdf2 renders tables best with explicit widths and borders
    html = re.sub(r"<table>", '<table border="1" width="100%">', html)

    pdf = _PDF(format="A4")
    pdf.footer_text = _footer_note(report)
    pdf.alias_nb_pages()
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("DejaVu", "", str(FONT_DIR / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(FONT_DIR / "DejaVuSans-Bold.ttf"))
    pdf.add_font("DejaVu", "I", str(FONT_DIR / "DejaVuSans-Oblique.ttf"))
    pdf.add_font("DejaVu", "BI", str(FONT_DIR / "DejaVuSans-BoldOblique.ttf"))
    pdf.add_font("DejaVuMono", "", str(FONT_DIR / "DejaVuSansMono.ttf"))
    pdf.add_page()
    pdf.set_font("DejaVu", size=10)
    ink = (27, 28, 31)
    pdf.write_html(
        html,
        font_family="DejaVu",
        warn_on_tags_not_matching=False,
        li_prefix_color=(31, 95, 191),
        tag_styles={
            "h1": TextStyle(font_family="DejaVu", font_style="B", font_size_pt=19, color=ink, t_margin=2, b_margin=3),
            "h2": TextStyle(font_family="DejaVu", font_style="B", font_size_pt=14, color=(31, 95, 191), t_margin=7, b_margin=2),
            "h3": TextStyle(font_family="DejaVu", font_style="B", font_size_pt=11.5, color=ink, t_margin=5, b_margin=1),
            "h4": TextStyle(font_family="DejaVu", font_style="B", font_size_pt=10.5, color=ink, t_margin=4, b_margin=1),
            "code": FontFace(family="DejaVuMono"),
            "pre": TextStyle(font_family="DejaVuMono", t_margin=3, b_margin=3),
            "a": FontFace(color=(31, 95, 191)),
        },
    )
    return bytes(pdf.output())


def to_markdown(report: dict) -> bytes:
    return (report["markdown"].rstrip() + f"\n\n<!-- {_footer_note(report)} -->\n").encode("utf-8")
